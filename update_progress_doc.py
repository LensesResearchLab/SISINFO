from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

source = r"C:\Users\jjhc7\OneDrive\Escritorio\Proyecto de Grado\Entregable_Avances_SISINFO.docx"
path = r"C:\Users\jjhc7\OneDrive\Escritorio\Proyecto de Grado\Entregable_Avances_SISINFO_Microsoft.docx"
doc = Document(source)

target = next((p for p in doc.paragraphs if p.text.strip() == "4. Próximos pasos"), None)
if target is None:
    raise RuntimeError("No se encontró la sección de próximos pasos")

def insert_before(text, style="Normal"):
    paragraph = OxmlElement("w:p")
    target._p.addprevious(paragraph)
    from docx.text.paragraph import Paragraph
    result = Paragraph(paragraph, target._parent)
    result.style = style
    result.add_run(text)
    return result

insert_before("3.5. Integración del inicio de sesión institucional", "Heading 2")
insert_before("Se implementó la integración del inicio de sesión de SISINFO con Microsoft Entra ID, el servicio de identidad institucional de Microsoft utilizado por la Universidad. El objetivo es que los usuarios ingresen con su cuenta institucional de Uniandes, sin depender de una contraseña almacenada exclusivamente en SISINFO.")
insert_before("La integración utiliza el protocolo OpenID Connect sobre OAuth 2.0 mediante el flujo de código de autorización para aplicaciones web. El usuario es enviado a Microsoft, realiza allí la autenticación y luego es redirigido nuevamente a SISINFO mediante una URI autorizada.")
insert_before("El flujo implementado funciona de la siguiente manera:")
for text in [
    "• El usuario selecciona «Iniciar sesión con Uniandes» en la pantalla de acceso.",
    "• SISINFO genera una solicitud de autenticación hacia Microsoft Entra ID.",
    "• Microsoft valida la identidad y las credenciales institucionales.",
    "• Microsoft redirige al usuario a la ruta de retorno configurada en SISINFO.",
    "• SISINFO valida la respuesta, obtiene el correo y el nombre institucional y crea la sesión local mediante una cookie segura.",
    "• El backend busca al usuario local; si no existe, lo crea automáticamente para evitar bloquear el primer acceso.",
]:
    insert_before(text)
insert_before("Para conectar cada ambiente se deben registrar en Microsoft Entra las siguientes URI de redirección:")
insert_before("• QA: https://qa-sisinfo.virtual.uniandes.edu.co/auth/microsoft/callback")
insert_before("• Producción: https://sisinfo.virtual.uniandes.edu.co/auth/microsoft/callback")
insert_before("En el entorno local se utiliza http://localhost:3000/auth/microsoft/callback. El identificador de aplicación, el identificador del tenant, el secreto de cliente y la URI de redirección se manejan mediante variables de entorno y no deben almacenarse directamente en el código fuente.")
insert_before("Como parte de la provisión automática, los usuarios nuevos se registran inicialmente en SISINFO. Las cuentas institucionales identificadas como profesores reciben el rol de profesor y las cuentas configuradas como coordinadores reciben el rol de coordinador. Los demás usuarios se provisionan como estudiantes de maestría para permitir el acceso inicial a la plataforma.")
insert_before("El secreto de cliente debe mantenerse únicamente en el servidor y debe revocarse y reemplazarse si se expone. La aplicación también utiliza una clave interna para proteger la comunicación entre el callback de autenticación y la API de SISINFO.")
insert_before("Estado del avance: la autenticación con Microsoft Entra ID fue probada localmente y permite regresar a SISINFO con una sesión válida. Queda pendiente validar las cuentas institucionales en los ambientes de QA y producción y confirmar con DSIT la asignación definitiva de roles y atributos académicos.")

doc.save(path)
print(path)
